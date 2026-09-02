#!/usr/bin/env python3
# Run: python3 generate-docx.py research.json [output.docx]

from __future__ import annotations

import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    print("Error: python-docx required. Run: pip3 install python-docx", file=sys.stderr)
    sys.exit(1)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def build_doc(data: dict, output_path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(f"Lesson Research: {data['passage']}", 0)

    meta = doc.add_paragraph()
    if data.get("teacher_name"):
        meta.add_run(f"Prepared For: {data['teacher_name']}\n")
    if data.get("group_name"):
        meta.add_run(f"Group: {data['group_name']}\n")
    if data.get("date"):
        meta.add_run(f"Date: {data['date']}")

    sections = [
        ("1. Passage Context", data["passage_context"]),
        ("2. Historical and Cultural Background", data["historical_background"]),
        ("4. Commentary Insights", data["commentary_insights"]),
    ]
    for title, body in sections:
        add_heading(doc, title, level=1)
        doc.add_paragraph(body)

    add_heading(doc, "3. Key Word Studies", level=1)
    table = doc.add_table(rows=1, cols=5)
    headers = [
        "English Word",
        "Transliteration",
        "Literal Meaning",
        "Range of Meaning",
        "Translation Comparison",
    ]
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for word in data["word_studies"]:
        row = table.add_row().cells
        row[0].text = word["english"]
        row[1].text = word["transliteration"]
        row[2].text = word["literal_meaning"]
        row[3].text = word["range_of_meaning"]
        row[4].text = word["translation_comparison"]

    add_heading(doc, "5. Cross-References and Parallel Passages", level=1)
    for item in data["cross_references"]:
        prefix = "[Read aloud] " if item.get("read_aloud") else ""
        doc.add_paragraph(
            f"{prefix}{item['reference']} ({item['connection_type']}) — {item['connection']}",
            style="List Bullet",
        )

    add_heading(doc, "6. Theological Themes", level=1)
    for theme in data["theological_themes"]:
        add_heading(doc, theme["title"], level=2)
        p = doc.add_paragraph()
        p.add_run("Textual Presence: ").bold = True
        p.add_run(theme["textual_presence"])
        p = doc.add_paragraph()
        p.add_run("Practical Implication: ").bold = True
        p.add_run(theme["practical_implication"])

    add_heading(doc, "7. Thinking Prompts and Discussion Starters", level=1)
    doc.add_paragraph("For you:").runs[0].bold = True
    for idx, prompt in enumerate(data["thinking_prompts"], start=1):
        doc.add_paragraph(f"{idx}. {prompt}")
    doc.add_paragraph("For the group:").runs[0].bold = True
    for idx, starter in enumerate(data["discussion_starters"], start=1):
        doc.add_paragraph(f"{idx}. {starter}")

    add_heading(doc, "8. Anticipated Questions", level=1)
    for item in data["anticipated_questions"]:
        p = doc.add_paragraph()
        p.add_run(f"\"{item['question']}\"").bold = True
        doc.add_paragraph(f"{item['answer']} ({item['status']})")

    doc.save(output_path)


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: generate-docx.py research.json [output.docx]", file=sys.stderr)
        sys.exit(1)

    json_path = Path(sys.argv[1])
    output_path = (
        Path(sys.argv[2])
        if len(sys.argv) > 2
        else json_path.with_suffix(".docx")
    )

    with json_path.open(encoding="utf-8") as handle:
        data = json.load(handle)

    build_doc(data, output_path)
    print(output_path)


if __name__ == "__main__":
    main()
