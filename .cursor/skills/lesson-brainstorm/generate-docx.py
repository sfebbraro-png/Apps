#!/usr/bin/env python3
"""Generate a lesson brief .docx from JSON."""

from __future__ import annotations

import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    print("Error: python-docx required. Run: pip install python-docx", file=sys.stderr)
    sys.exit(1)


def add_labeled_paragraph(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    run_label = paragraph.add_run(f"{label}: ")
    run_label.bold = True
    paragraph.add_run(value)


def build_doc(data: dict, output_path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading("Lesson Brief", 0)

    meta_lines = [
        ("Passage", data["passage"]),
        ("Series", data.get("series")),
        ("Date", data.get("date")),
        ("Teacher", data.get("teacher_name")),
    ]
    for label, value in meta_lines:
        if value:
            add_labeled_paragraph(doc, label, value)

    doc.add_paragraph()

    sections = [
        ("Big Idea", data["big_idea"]),
        ("Key Tension", data["key_tension"]),
        ("Group Need", data["group_need"]),
        ("Desired Response", data["desired_response"]),
        ("The Turn", data["the_turn"]),
        ("Opening Question", data["opening_question"]),
        ("Likely Sticking Point", data["sticking_point"]),
        ("One Image or Illustration Idea", data["illustration_idea"]),
    ]
    for label, value in sections:
        doc.add_heading(label, level=2)
        doc.add_paragraph(value)

    doc.add_heading("Supporting Passages", level=2)
    for item in data["supporting_passages"]:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{item['reference']} — ").bold = True
        paragraph.add_run(item["note"])

    footer = doc.add_paragraph()
    footer.add_run(
        "This brief is a launchpad, not a script. Take it to prayer and make it yours."
    ).italic = True

    doc.save(output_path)


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: generate-docx.py brief.json [output.docx]", file=sys.stderr)
        sys.exit(1)

    json_path = Path(sys.argv[1])
    output_path = (
        Path(sys.argv[2])
        if len(sys.argv) > 2
        else json_path.with_suffix(".docx")
    )

    with json_path.open(encoding="utf-8") as handle:
        data = json.load(handle)

    build_doc(data, output_path)
    print(output_path)


if __name__ == "__main__":
    main()
